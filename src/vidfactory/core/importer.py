"""One-time import of the legacy Flugbuch.xlsx flight log into SQLite.

Seeds the Site master (launch/landing + elevations from the DropDownData sheet) and the 598 Outings
from the Flugbuch sheet. Idempotent via `replace=True` (clears existing outings/sites first), so it
can be re-run safely while there are no video projects yet.

CLI:  python -m vidfactory.core.importer /path/to/Flugbuch.xlsx
"""

from __future__ import annotations

import datetime
import logging
from pathlib import Path

import openpyxl
from sqlalchemy import delete
from sqlalchemy.orm import Session

from vidfactory.database.models import Outing, Project, Site, outing_buddies

logger = logging.getLogger(__name__)


def _s(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _i(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None


def _f(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _launch_type(value) -> str | None:
    """Normalise the legacy f/r codes to the readable words used in the app."""
    text = _s(value)
    if text is None:
        return None
    return {"f": "forward", "r": "reverse"}.get(text.lower(), text)


def _load_site_master(ws) -> list[tuple[str, str, int | None]]:
    """Launch in cols A/B, landing in cols D/E, from row 3 (headers in rows 1-2)."""
    out: list[tuple[str, str, int | None]] = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        launch_name, launch_h = _s(row[0]), _i(row[1])
        land_name, land_h = _s(row[3]), _i(row[4])
        if launch_name:
            out.append((launch_name, "launch", launch_h))
        if land_name:
            out.append((land_name, "landing", land_h))
    return out


def import_flugbuch(path: str | Path, db: Session, replace: bool = True) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    if replace:
        db.execute(outing_buddies.delete())
        db.execute(delete(Project))
        db.execute(delete(Outing))
        db.execute(delete(Site))
        db.flush()

    site_map: dict[tuple[str, str], int] = {}

    def site_id(name, kind: str, elevation) -> int | None:
        name = _s(name)
        if not name:
            return None
        key = (name, kind)
        if key not in site_map:
            site = Site(name=name, kind=kind, elevation_m=_i(elevation))
            db.add(site)
            db.flush()
            site_map[key] = site.id
        return site_map[key]

    # 1) Site master from DropDownData (canonical elevations).
    for name, kind, elevation in _load_site_master(wb["DropDownData"]):
        site_id(name, kind, elevation)
    master_sites = len(site_map)

    # 2) Outings from the Flugbuch sheet (row 1 is the header).
    outings = 0
    for row in wb["Flugbuch"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        date = row[0]
        if isinstance(date, datetime.datetime):
            date = date.date()
        db.add(
            Outing(
                date=date,
                launch_site_id=site_id(row[1], "launch", row[2]),
                landing_site_id=site_id(row[3], "landing", row[4]),
                flight_time_min=_i(row[6]),
                distance_km=_f(row[7]),
                category=_s(row[8]),
                max_alt_m=_i(row[9]),
                glider=_s(row[11]),
                harness=_s(row[12]),
                launch_type=_launch_type(row[13]),
                comment=_s(row[14]),
            )
        )
        outings += 1

    db.commit()

    # Keep the form dropdowns in sync with whatever the import introduced.
    from vidfactory.core import lookups

    lookups.sync_from_outings(db)

    summary = {
        "sites_total": len(site_map),
        "sites_from_master": master_sites,
        "sites_added_from_outings": len(site_map) - master_sites,
        "outings": outings,
    }
    logger.info("Flugbuch import: %s", summary)
    return summary


def _main() -> None:
    import sys

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    if len(sys.argv) != 2:
        print("usage: python -m vidfactory.core.importer /path/to/Flugbuch.xlsx")
        raise SystemExit(2)

    from vidfactory.database.db import get_engine, init_db

    init_db()
    with Session(get_engine()) as db:
        print(import_flugbuch(sys.argv[1], db))


if __name__ == "__main__":
    _main()
